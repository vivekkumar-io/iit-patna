"""
Generate submission PowerPoint for Enterprise Knowledge Assistant.

Author: Vivek Kumar
"""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

PROJECT_ROOT = Path(__file__).resolve().parent.parent
OUTPUT_PATH = PROJECT_ROOT / "SubmissionDocument" / "Enterprise_Knowledge_Assistant_Vivek_kumar.pptx"
OUTPUT_PATH_FALLBACK = (
    PROJECT_ROOT / "SubmissionDocument" / "Enterprise_Knowledge_Assistant_Vivek_kumar_updated.pptx"
)
SCREENSHOT_PATH = (
    PROJECT_ROOT / "SubmissionDocument" / "screenshots" / "app_home_screen.png"
)
ARCHITECTURE_DIAGRAM_PATH = (
    PROJECT_ROOT
    / "SubmissionDocument"
    / "screenshots"
    / "architecture_flow_diagram.png"
)
ARCHITECTURE_DIAGRAM_FALLBACK = (
    Path(r"D:\_GenAIProject\BKp\Enterprise_Knowledge_Assistant_temp")
    / "screenshots"
    / "architecture_flow_diagram.png"
)

DARK_BLUE = RGBColor(0x15, 0x65, 0xC0)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DARK_TEXT = RGBColor(0x1F, 0x29, 0x37)
GRAY = RGBColor(0x64, 0x74, 0x8B)

ASCII_ARCHITECTURE_DIAGRAM = """Documents (PDF/TXT/DOCX)
        |
        v
+-------------------+
|  Document Loader  |  ingestion/loader.py
+---------+---------+
          v
+-------------------+
|     Chunker       |  ingestion/chunker.py
+---------+---------+
          v
+-------------------+
|   Embeddings      |  ingestion/embedder.py
+---------+---------+
          +----------------------+
          v                      v
+-------------------+    +-------------------+
|  ChromaDB Vector  |    |    BM25 Index     |
|      Store        |    |                   |
+---------+---------+    +---------+---------+
          |                        |
          +------------+-----------+
                       v
            +-----------------------+
            |   Hybrid Retriever    |  retrieval/hybrid_retriever.py
            |   (Vector + BM25)     |
            +-----------+-----------+
                        v
            +-----------------------+
            |      Reranker         |  retrieval/reranker.py
            +-----------+-----------+
                        v
            +-----------------------+
            |   Context Builder     |  generation/guard.py
            +-----------+-----------+
                        v
            +-----------------------+
            |  LLM + Chat Memory    |  generation/chain.py
            +-----------+-----------+
                        v
            +-----------------------+
            |   Streamlit Chat UI   |  app/ui/chat.py
            |  Answer + Sources     |
            +-----------------------+"""


def resolve_architecture_diagram_path() -> Path:
    if ARCHITECTURE_DIAGRAM_PATH.exists():
        return ARCHITECTURE_DIAGRAM_PATH
    if ARCHITECTURE_DIAGRAM_FALLBACK.exists():
        return ARCHITECTURE_DIAGRAM_FALLBACK
    return ARCHITECTURE_DIAGRAM_PATH


def set_title_style(shape, size=32, color=DARK_BLUE, bold=True):
    shape.text_frame.paragraphs[0].font.size = Pt(size)
    shape.text_frame.paragraphs[0].font.bold = bold
    shape.text_frame.paragraphs[0].font.color.rgb = color


def add_bullets(text_frame, items, size=18, color=DARK_TEXT, spacing=8):
    text_frame.clear()
    for index, item in enumerate(items):
        paragraph = text_frame.paragraphs[0] if index == 0 else text_frame.add_paragraph()
        paragraph.text = item
        paragraph.level = 0
        paragraph.font.size = Pt(size)
        paragraph.font.color.rgb = color
        paragraph.space_after = Pt(spacing)


def add_title_slide(prs, title, subtitle):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = DARK_BLUE

    title_box = slide.shapes.add_textbox(Inches(0.7), Inches(2.0), Inches(8.6), Inches(1.5))
    title_frame = title_box.text_frame
    title_frame.text = title
    title_para = title_frame.paragraphs[0]
    title_para.font.size = Pt(36)
    title_para.font.bold = True
    title_para.font.color.rgb = WHITE
    title_para.alignment = PP_ALIGN.CENTER

    subtitle_box = slide.shapes.add_textbox(Inches(0.7), Inches(3.6), Inches(8.6), Inches(1.2))
    subtitle_frame = subtitle_box.text_frame
    subtitle_frame.text = subtitle
    subtitle_para = subtitle_frame.paragraphs[0]
    subtitle_para.font.size = Pt(20)
    subtitle_para.font.color.rgb = WHITE
    subtitle_para.alignment = PP_ALIGN.CENTER


def add_content_slide(prs, title, bullets):
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.35), Inches(9), Inches(0.8))
    title_frame = title_box.text_frame
    title_frame.text = title
    set_title_style(title_box, size=30)

    body_box = slide.shapes.add_textbox(Inches(0.7), Inches(1.2), Inches(8.8), Inches(5.8))
    add_bullets(body_box.text_frame, bullets)


def add_ui_screenshot_slide(prs, title, image_path, bullets):
    """Slide with application screenshot and UI functionality notes."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.25), Inches(9), Inches(0.6))
    title_box.text_frame.text = title
    set_title_style(title_box, size=28)

    if image_path.exists():
        slide.shapes.add_picture(
            str(image_path),
            Inches(0.45),
            Inches(0.95),
            width=Inches(5.9),
        )
    else:
        placeholder = slide.shapes.add_textbox(Inches(0.45), Inches(2.5), Inches(5.9), Inches(1))
        placeholder.text_frame.text = "Screenshot not found"
        placeholder.text_frame.paragraphs[0].font.size = Pt(14)

    body_box = slide.shapes.add_textbox(Inches(6.5), Inches(0.95), Inches(3.2), Inches(6.2))
    add_bullets(body_box.text_frame, bullets, size=13, spacing=6)


def add_image_slide(prs, title, image_path, caption=None):
    """Full-width slide for architecture or diagram images."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.2), Inches(9), Inches(0.55))
    title_box.text_frame.text = title
    set_title_style(title_box, size=28)

    if image_path.exists():
        slide.shapes.add_picture(
            str(image_path),
            Inches(0.35),
            Inches(0.85),
            width=Inches(9.3),
        )
    else:
        placeholder = slide.shapes.add_textbox(Inches(0.5), Inches(3.0), Inches(9), Inches(1))
        placeholder.text_frame.text = "Image not found"
        placeholder.text_frame.paragraphs[0].font.size = Pt(14)

    if caption:
        caption_box = slide.shapes.add_textbox(Inches(0.5), Inches(7.05), Inches(9), Inches(0.35))
        caption_box.text_frame.text = caption
        caption_box.text_frame.paragraphs[0].font.size = Pt(11)
        caption_box.text_frame.paragraphs[0].font.color.rgb = GRAY


def add_text_slide(prs, title, body_text):
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.35), Inches(9), Inches(0.8))
    title_box.text_frame.text = title
    set_title_style(title_box, size=30)

    body_box = slide.shapes.add_textbox(Inches(0.7), Inches(1.2), Inches(8.8), Inches(5.8))
    body_frame = body_box.text_frame
    body_frame.word_wrap = True
    body_frame.text = body_text
    for paragraph in body_frame.paragraphs:
        paragraph.font.size = Pt(16)
        paragraph.font.color.rgb = DARK_TEXT
        paragraph.space_after = Pt(10)


def add_ascii_diagram_slide(prs, title, diagram_text):
    """Monospace slide for the high-level ASCII architecture from ARCHITECTURE.md."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.15), Inches(9), Inches(0.5))
    title_box.text_frame.text = title
    set_title_style(title_box, size=26)

    body_box = slide.shapes.add_textbox(Inches(0.25), Inches(0.65), Inches(9.5), Inches(6.7))
    body_frame = body_box.text_frame
    body_frame.word_wrap = False
    body_frame.text = diagram_text
    for paragraph in body_frame.paragraphs:
        paragraph.font.name = "Courier New"
        paragraph.font.size = Pt(7.5)
        paragraph.font.color.rgb = DARK_TEXT
        paragraph.space_after = Pt(0)
        paragraph.line_spacing = 1.0


def build_presentation() -> Path:
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)

    add_title_slide(
        prs,
        "Enterprise Knowledge Assistant",
        "Project 2 — Advanced RAG Application\n"
        "Vivek Kumar | IIT Patna GenAI Development Program\n"
        "Oeeggis Corporation (Fictional — Academic Project)",
    )

    add_content_slide(
        prs,
        "Problem Statement",
        [
            "Employees need quick answers from HR, IT, travel, and benefits policies.",
            "Policy information is spread across multiple PDF and text documents.",
            "Manual search is slow, inconsistent, and time-consuming.",
            "Need: a reliable assistant that answers from official company documents only.",
        ],
    )

    add_content_slide(
        prs,
        "Solution Overview",
        [
            "Built an Employee Knowledge Assistant using Retrieval-Augmented Generation (RAG).",
            "Indexes local company documents and retrieves relevant chunks for each question.",
            "Uses hybrid search + reranking for better accuracy.",
            "Generates grounded answers with source citations in a Streamlit chat UI.",
            "Remembers follow-up questions using conversational memory.",
        ],
    )

    add_image_slide(
        prs,
        "Architecture Flow Diagram",
        resolve_architecture_diagram_path(),
        "A: Indexing Pipeline (offline)  |  B: Query Pipeline (online)",
    )

    add_ascii_diagram_slide(
        prs,
        "High-Level Architecture (ASCII)",
        ASCII_ARCHITECTURE_DIAGRAM,
    )

    add_content_slide(
        prs,
        "Architecture Summary",
        [
            "Indexing: Documents → Loader → Chunker → Embeddings → ChromaDB + BM25.",
            "Query: User input → route → memory rewrite → hybrid search → rerank → context builder → LLM.",
            "Routing: greetings (UI); casual intros and memory questions get instant replies without retrieval.",
            "Policy questions run the full RAG pipeline and return answers with source citations.",
            "Hallucination guard returns a safe not-found message when no relevant chunks are found.",
            "Streamlit UI shows chat history, sources, Knowledge Scope filter, and Clear/Reset.",
        ],
    )

    add_content_slide(
        prs,
        "Technology Stack",
        [
            "Language: Python 3.10+",
            "UI: Streamlit",
            "Framework: LangChain",
            "LLM & Embeddings: OpenAI (gpt-4o-mini, text-embedding-3-small)",
            "Vector Database: ChromaDB",
            "Keyword Search: BM25 (rank-bm25)",
            "Reranking: sentence-transformers cross-encoder",
            "Config: pydantic-settings + .env",
        ],
    )

    add_content_slide(
        prs,
        "Key Features Implemented",
        [
            "Document ingestion from local folder (PDF, DOCX, TXT)",
            "Chunking, embeddings, and vector storage",
            "Semantic retrieval + BM25 hybrid search",
            "Cross-encoder reranking for relevance",
            "Conversational memory for follow-up questions",
            "Source citations displayed in the UI",
            "Hallucination guard with clear not-found responses",
            "Startup preload of search models",
        ],
    )

    add_content_slide(
        prs,
        "Sample Documents",
        [
            "Leave_Policy.pdf",
            "HR_Handbook.pdf",
            "IT_Policy.txt",
            "Travel_Policy.txt",
            "Benefits_Documentation.txt",
            "Code_of_Conduct.txt",
            "Company_FAQs.txt",
            "Evaluator can run the app without creating data from scratch.",
        ],
    )

    add_ui_screenshot_slide(
        prs,
        "Application UI — Home Screen",
        SCREENSHOT_PATH,
        [
            "Header: Oeeggis Corporation branding and app title.",
            "Sidebar — Knowledge Scope lists all indexed policy documents.",
            "User can see which topics the assistant can answer from.",
            "Main chat area shows welcome message from the assistant.",
            "Chat input at bottom: type policy questions in natural language.",
            "Clear / Reset button clears chat history and memory.",
            "Footer disclaimer: academic/fictional company notice.",
            "On startup, search models preload in the background.",
        ],
    )

    add_content_slide(
        prs,
        "UI Functionality Explained",
        [
            "Knowledge Scope (sidebar): all 7 documents are searchable.",
            "Chat interaction: ask HR, IT, travel, leave, and benefits questions.",
            "Conversation history: previous messages stay visible in the chat.",
            "Source citations: policy answers show document names below reply.",
            "Follow-up support: e.g. 'What about carry-forward?' after leave policy.",
            "Intro handling: casual messages get quick replies without policy search.",
            "Hallucination guard: returns not-found message when data is missing.",
            "Clear / Reset: starts a fresh session for a new conversation.",
        ],
    )

    add_content_slide(
        prs,
        "Demo Flow",
        [
            "1. Run setup.bat and python scripts/ingest.py",
            "2. Start app with start.bat → http://localhost:8501",
            "3. Ask: What is the leave policy?",
            "4. Follow-up: What about carry-forward?",
            "5. Ask: What is the VPN policy?",
            "6. Test not-found: What is the salary for the CEO?",
            "7. Show source citations and Clear/Reset button",
        ],
    )

    add_content_slide(
        prs,
        "Sample Output",
        [
            "Q: What is the leave policy?",
            "A: Summary from leave policy documents.",
            "Sources: Leave_Policy.pdf",
            "",
            "Q: What is the salary for the CEO?",
            "A: I could not find this information in the available documents.",
            "",
            "Q: Hi, my name is Abhishek from Bangalore.",
            "A: Brief acknowledgement without unsolicited policy advice.",
        ],
    )

    add_content_slide(
        prs,
        "Key Design Decisions",
        [
            "Hybrid retrieval combines semantic meaning and exact keyword matches.",
            "Reranking improves precision before sending context to the LLM.",
            "Conversational routing skips document search for casual introductions.",
            "Models preload at startup to reduce first-question delay.",
            "Modular layers: ingestion, retrieval, generation, and app UI.",
        ],
    )

    add_content_slide(
        prs,
        "Limitations",
        [
            "Requires OpenAI API key (paid usage for embeddings and chat).",
            "First startup may take 60–90 seconds while models load.",
            "Answers limited to indexed documents only.",
            "No user authentication or role-based access control.",
            "Session-based memory (not persisted across server restarts).",
        ],
    )

    add_content_slide(
        prs,
        "How to Run",
        [
            "1. setup.bat (first time only)",
            "2. copy .env.example .env and add OPENAI_API_KEY",
            "3. python scripts/ingest.py",
            "4. start.bat",
            "5. Open http://localhost:8501",
            "Submission docs available in SubmissionDocument/ folder.",
        ],
    )

    add_title_slide(
        prs,
        "Thank You",
        "Questions & Discussion\n\nVivek Kumar\nEnterprise Knowledge Assistant — Advanced RAG",
    )

    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    try:
        prs.save(OUTPUT_PATH)
        return OUTPUT_PATH
    except PermissionError:
        prs.save(OUTPUT_PATH_FALLBACK)
        return OUTPUT_PATH_FALLBACK


if __name__ == "__main__":
    path = build_presentation()
    print(f"Presentation created: {path}")
