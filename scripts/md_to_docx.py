"""
Convert markdown files to Word (.docx) for evaluators who cannot open .md files.

Author: Vivek Kumar
"""

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.shared import Pt, RGBColor

PROJECT_ROOT = Path(__file__).resolve().parent.parent
SUBMISSION_DIR = PROJECT_ROOT / "SubmissionDocument"

FILES = (
    "ARCHITECTURE.md",
    "ENVIRONMENT_SETUP.md",
    "SAMPLE_OUTPUT.md",
)


def set_run_code_style(run) -> None:
    run.font.name = "Courier New"
    run.font.size = Pt(9)


def add_formatted_text(paragraph, text: str) -> None:
    """Add text with basic **bold** and `inline code` support."""
    pattern = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`)")
    pos = 0
    for match in pattern.finditer(text):
        if match.start() > pos:
            paragraph.add_run(text[pos : match.start()])
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        else:
            run = paragraph.add_run(token[1:-1])
            set_run_code_style(run)
        pos = match.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])


def add_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    col_count = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=col_count)
    table.style = "Table Grid"
    for row_idx, row in enumerate(rows):
        for col_idx in range(col_count):
            cell_text = row[col_idx] if col_idx < len(row) else ""
            cell = table.rows[row_idx].cells[col_idx]
            cell.text = ""
            paragraph = cell.paragraphs[0]
            add_formatted_text(paragraph, cell_text.strip())
            for run in paragraph.runs:
                run.font.size = Pt(10)


def convert_markdown_to_docx(md_path: Path, docx_path: Path) -> None:
    lines = md_path.read_text(encoding="utf-8").splitlines()
    doc = Document()

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    in_code_block = False
    code_lines: list[str] = []
    table_rows: list[list[str]] = []
    in_table = False

    for line in lines:
        stripped = line.rstrip()

        if stripped.startswith("```"):
            if in_code_block:
                paragraph = doc.add_paragraph()
                run = paragraph.add_run("\n".join(code_lines))
                set_run_code_style(run)
                run.font.size = Pt(9)
                code_lines = []
                in_code_block = False
            else:
                in_code_block = True
            continue

        if in_code_block:
            code_lines.append(line)
            continue

        if "|" in stripped and stripped.strip().startswith("|"):
            cells = [cell.strip() for cell in stripped.strip("|").split("|")]
            if all(set(cell) <= {"-", ":"} for cell in cells):
                continue
            if not in_table:
                in_table = True
                table_rows = []
            table_rows.append(cells)
            continue

        if in_table:
            add_table(doc, table_rows)
            table_rows = []
            in_table = False

        if not stripped:
            continue

        if stripped == "---":
            doc.add_paragraph()
            continue

        if stripped.startswith("# "):
            heading = doc.add_heading(stripped[2:].strip(), level=0)
            heading.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT
            continue

        if stripped.startswith("## "):
            doc.add_heading(stripped[3:].strip(), level=1)
            continue

        if stripped.startswith("### "):
            doc.add_heading(stripped[4:].strip(), level=2)
            continue

        if stripped.startswith("- "):
            paragraph = doc.add_paragraph(style="List Bullet")
            add_formatted_text(paragraph, stripped[2:].strip())
            continue

        numbered = re.match(r"^(\d+)\.\s+(.*)$", stripped)
        if numbered:
            paragraph = doc.add_paragraph(style="List Number")
            add_formatted_text(paragraph, numbered.group(2).strip())
            continue

        paragraph = doc.add_paragraph()
        add_formatted_text(paragraph, stripped)

    if in_table:
        add_table(doc, table_rows)

    if in_code_block and code_lines:
        paragraph = doc.add_paragraph()
        run = paragraph.add_run("\n".join(code_lines))
        set_run_code_style(run)

    docx_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(docx_path)


def main() -> None:
    targets = sys.argv[1:] if len(sys.argv) > 1 else FILES
    for name in targets:
        md_path = SUBMISSION_DIR / name
        if not md_path.exists():
            md_path = PROJECT_ROOT / name
        if not md_path.exists():
            print(f"Skip (not found): {name}")
            continue
        docx_path = md_path.with_suffix(".docx")
        convert_markdown_to_docx(md_path, docx_path)
        print(f"Created: {docx_path}")


if __name__ == "__main__":
    main()
